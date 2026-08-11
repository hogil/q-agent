#!/usr/bin/env python3
"""AX_Award_지원서.html -> 제출용 .docx (A4 5장, 화면과 동일).

왜 이미지인가
  본문 시안이 flex·grid·border-radius 로 짜여 있어 Word 표로는 재현이 안 된다.
  페이지를 240dpi 로 떠서 여백 0 의 A4 에 full-bleed 로 앉히면 화면과 1:1 로 같다.
  대신 텍스트 선택·편집은 안 된다 — 편집이 필요하면 HTML 을 고치고 다시 뽑는다.

쓰기
  python build_docx.py            # 240dpi
  python build_docx.py 300        # dpi 지정
"""
import sys, pathlib, asyncio
from playwright.async_api import async_playwright
from docx import Document
from docx.shared import Cm, Pt, Emu
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

HERE = pathlib.Path(__file__).parent
SRC = HERE / "AX_Award_지원서.html"
PNGDIR = HERE / "docx_pages"
OUT = HERE / "AX_Award_지원서.docx"

DPI = int(sys.argv[1]) if len(sys.argv) > 1 else 240
PAGE_W_CM, PAGE_H_CM = 21.0, 29.7
CSS_PX_W = 794                      # A4 폭 @96dpi
SCALE = DPI / 96                    # 브라우저 device_scale_factor


async def shoot():
    PNGDIR.mkdir(exist_ok=True)
    async with async_playwright() as pw:
        b = await pw.chromium.launch()
        pg = await b.new_page(viewport={"width": 900, "height": 1200},
                              device_scale_factor=SCALE)
        await pg.goto(SRC.as_uri())
        await pg.wait_for_timeout(800)          # 내장 폰트 로드 대기
        els = await pg.query_selector_all("section.page")
        paths = []
        for i, el in enumerate(els, 1):
            p = PNGDIR / f"p{i}.png"
            await el.screenshot(path=str(p))
            paths.append(p)
        await b.close()
        return paths


def build(paths):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(0)
    s.header_distance = s.footer_distance = Cm(0)

    for i, p in enumerate(paths):
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        # 줄 높이를 그림 높이에 정확히 맞춘다 — 안 그러면 leading 이 붙어 다음 장으로 흐른다
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Emu(Cm(PAGE_H_CM).emu)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if i:
            pf.page_break_before = True
        para.add_run().add_picture(str(p), width=Cm(PAGE_W_CM), height=Cm(PAGE_H_CM))

    doc.save(OUT)


paths = asyncio.run(shoot())
build(paths)
mb = OUT.stat().st_size / 1e6
print(f"{DPI}dpi · {len(paths)}장 -> {OUT}  ({mb:.1f} MB)")
